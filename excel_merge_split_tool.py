import sys
import os
import traceback
from datetime import datetime
import pandas as pd
from PyQt5 import QtWidgets
from PyQt5.QtWidgets import (
    QApplication,
    QMainWindow,
    QWidget,
    QFileDialog,
    QMessageBox,
    QVBoxLayout,
    QHBoxLayout,
    QLabel,
    QPushButton,
    QListWidget,
    QCheckBox,
    QComboBox,
    QProgressBar,
    QRadioButton,
    QTextEdit,
    QGroupBox,
)

# For writing styles and fixing column width
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

SUPPORTED_EXTENSIONS = (".xlsx", ".xls", ".csv")

# ---------- Utility functions ----------
def list_excel_files_in_folder(folder):
    """
    Recursively lists all supported Excel and CSV files in a given folder.
    """
    out = []
    for root, _, files in os.walk(folder):
        for f in files:
            if f.lower().endswith(SUPPORTED_EXTENSIONS):
                out.append(os.path.join(root, f))
    return out

def read_file_to_df(path):
    """
    Reads a single supported file (Excel or CSV) into a pandas DataFrame.
    For Excel files, it reads only the first sheet.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        return pd.read_csv(path, dtype=str, keep_default_na=False)
    else:
        try:
            return pd.read_excel(path, sheet_name=0, dtype=str, keep_default_na=False)
        except Exception:
            return pd.read_excel(
                path, sheet_name=0, engine="xlrd", dtype=str, keep_default_na=False
            )

def read_all_sheets(path):
    """
    Reads all sheets from a supported file into a dictionary of DataFrames.
    Return dict of sheetname -> DataFrame.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        return {"Sheet1": pd.read_csv(path, dtype=str, keep_default_na=False)}
    else:
        return pd.read_excel(path, sheet_name=None, dtype=str, keep_default_na=False)

def save_df_to_excel(
    df, path, sheet_name="Sheet1", number_format_map=None, apply_theme=False
):
    """
    Save a single DataFrame to an .xlsx file using openpyxl.
    """
    wb = Workbook()
    ws = wb.active
    if ws is not None:
        ws.title = sheet_name
    else:
        ws = wb.create_sheet(title=sheet_name)

    headers = list(df.columns)
    ws.append(headers)
    for _, row in df.iterrows():
        values = [row.get(c) for c in headers]
        ws.append(values)

    thin = Side(border_style="thin", color="000000")
    for col_index, col in enumerate(headers, start=1):
        letter = get_column_letter(col_index)
        cell = ws[f"{letter}1"]
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    if apply_theme:
        pass

    for col in ws.columns:
        max_len = 0
        col_index = col[0].column
        if col_index is None or not isinstance(col_index, int):
            continue
        col_letter = get_column_letter(col_index)
        for cell in col:
            try:
                val = cell.value
                length = len(str(val)) if val is not None else 0
            except:
                length = 0
            if length > max_len:
                max_len = length
        ws.column_dimensions[col_letter].width = min(max(50, max_len + 2), 100)

    output_directory = os.path.dirname(path)
    if output_directory and not os.path.exists(output_directory):
        os.makedirs(output_directory)
    wb.save(path)

def apply_openpyxl_autofit_and_theme(
    path, sheet_name=None, number_format_map=None, apply_theme=False
):
    """
    Post-process existing workbook to set column widths.
    """
    wb = load_workbook(path)
    if sheet_name:
        if sheet_name in wb.sheetnames:
            ws_list = [wb[sheet_name]]
        else:
            ws_list = []
    else:
        ws_list = [wb[s] for s in wb.sheetnames]

    thin = Side(border_style="thin", color="000000")
    for ws in ws_list:
        if ws is None: continue
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="DDDDDD")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    val = cell.value
                    length = len(str(val)) if val is not None else 0
                except:
                    length = 0
                if length > max_len:
                    max_len = length
            ws.column_dimensions[col_letter].width = min(max(50, max_len + 2), 100)

    if apply_theme:
        pass
    wb.save(path)

# ---------- GUI ----------
class ExcelMergeSplitWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Excel Merge/Split Tool")
        self.resize(800, 600)
        self.selected_files = []
        self.output_folder = None
        self.overwrite_originals = False # Not used for Merge/Split mostly, but good for Split

        self._build_ui()

    def _build_ui(self):
        main = QWidget()
        main_layout = QVBoxLayout()
        main.setLayout(main_layout)
        self.setCentralWidget(main)

        # A. File Handling
        file_group = QGroupBox("1. File Selection")
        fg_layout = QVBoxLayout()
        file_group.setLayout(fg_layout)

        self.file_list_widget = QListWidget()
        fg_layout.addWidget(self.file_list_widget)

        btn_row = QWidget()
        br_layout = QHBoxLayout()
        btn_row.setLayout(br_layout)
        br_layout.setContentsMargins(0, 0, 0, 0)
        self.add_files_btn = QPushButton("Add Files")
        self.add_folder_btn = QPushButton("Add Folder")
        self.clear_files_btn = QPushButton("Clear")
        br_layout.addWidget(self.add_files_btn)
        br_layout.addWidget(self.add_folder_btn)
        br_layout.addWidget(self.clear_files_btn)
        fg_layout.addWidget(btn_row)

        self.add_files_btn.clicked.connect(self.add_files)
        self.add_folder_btn.clicked.connect(self.add_folder)
        self.clear_files_btn.clicked.connect(self.clear_files)

        # Output folder
        out_row = QWidget()
        or_layout = QHBoxLayout()
        out_row.setLayout(or_layout)
        or_layout.setContentsMargins(0, 0, 0, 0)
        self.out_folder_label = QLabel("Output folder: (not set)")
        self.select_out_btn = QPushButton("Select Output Folder")
        or_layout.addWidget(self.out_folder_label)
        or_layout.addWidget(self.select_out_btn)
        fg_layout.addWidget(out_row)
        self.select_out_btn.clicked.connect(self.select_output_folder)

        # B. Tasks
        tasks_group = QGroupBox("2. Select Operation")
        tasks_layout = QVBoxLayout()
        tasks_group.setLayout(tasks_layout)

        self.chk_merge = QCheckBox("Merge selected files into one workbook")
        merge_row = QWidget()
        mr_layout = QHBoxLayout()
        merge_row.setLayout(mr_layout)
        mr_layout.setContentsMargins(0, 0, 0, 0)
        mr_layout.addWidget(QLabel("   Merged filename:"))
        self.merge_name_edit = QtWidgets.QLineEdit("merged_output.xlsx")
        mr_layout.addWidget(self.merge_name_edit)
        tasks_layout.addWidget(self.chk_merge)
        tasks_layout.addWidget(merge_row)

        self.chk_split = QCheckBox("Split workbook sheets into separate files")
        tasks_layout.addWidget(self.chk_split)

        # C. Process
        process_group = QGroupBox("3. Process")
        process_layout = QVBoxLayout()
        process_group.setLayout(process_layout)

        self.process_btn = QPushButton("Start Processing")
        self.process_btn.clicked.connect(self.start_processing)
        process_layout.addWidget(self.process_btn)

        self.progress = QProgressBar()
        self.progress.setValue(0)
        process_layout.addWidget(self.progress)

        # Log
        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)

        main_layout.addWidget(file_group)
        main_layout.addWidget(tasks_group)
        main_layout.addWidget(process_group)
        main_layout.addWidget(QLabel("Log / Report"))
        main_layout.addWidget(self.log_text)

    # ---------- UI Actions ----------
    def add_files(self):
        files, _ = QFileDialog.getOpenFileNames(
            self, "Select Excel/CSV files", "", "Spreadsheet files (*.xlsx *.xls *.csv)"
        )
        if files:
            for f in files:
                if f not in self.selected_files:
                    self.selected_files.append(f)
                    self.file_list_widget.addItem(f)
            self.log(f"Added {len(files)} files.")

    def add_folder(self):
        folder = QFileDialog.getExistingDirectory(
            self, "Select folder containing Excel/CSV files", ""
        )
        if folder:
            files = list_excel_files_in_folder(folder)
            added = 0
            for f in files:
                if f not in self.selected_files:
                    self.selected_files.append(f)
                    self.file_list_widget.addItem(f)
                    added += 1
            self.log(f"Added {added} files from folder {folder}.")

    def clear_files(self):
        self.selected_files = []
        self.file_list_widget.clear()
        self.log("Cleared file list.")

    def select_output_folder(self):
        folder = QFileDialog.getExistingDirectory(self, "Select output folder", "")
        if folder:
            self.output_folder = folder
            self.out_folder_label.setText(f"Output folder: {folder}")
            self.log(f"Output folder set: {folder}")

    def start_processing(self):
        if not self.file_list_widget.count():
            QMessageBox.warning(self, "No files", "No files selected.")
            return

        if not self.chk_merge.isChecked() and not self.chk_split.isChecked():
            QMessageBox.warning(self, "No operation", "Please select Merge or Split operation.")
            return

        if not self.output_folder:
             QMessageBox.warning(self, "No Output Folder", "Please select an output folder.")
             return

        files = [self.file_list_widget.item(i).text() for i in range(self.file_list_widget.count())]

        # Merge Logic
        if self.chk_merge.isChecked():
            merged_name = self.merge_name_edit.text().strip() or "merged_output.xlsx"
            merged_path = os.path.join(self.output_folder, merged_name)
            self.log("Starting merge...")
            try:
                self.merge_files(files, merged_path)
                self.log(f"Merged output saved: {merged_path}")
            except Exception as e:
                self.log(f"Merge failed: {e}\n{traceback.format_exc()}")

        # Split Logic
        if self.chk_split.isChecked():
            total = len(files)
            self.progress.setMaximum(total)
            self.progress.setValue(0)
            for idx, file_path in enumerate(files, start=1):
                self.log(f"Splitting: {file_path}")
                try:
                    sheets = read_all_sheets(file_path)
                    base_name = os.path.splitext(os.path.basename(file_path))[0]
                    for sheetname, df in sheets.items():
                        out_name = f"{base_name}__{sheetname[:20]}.xlsx"
                        out_path = os.path.join(self.output_folder, out_name)
                        save_df_to_excel(df, out_path, sheet_name=sheetname)
                        self.log(f"  Saved sheet: {out_path}")
                except Exception as e:
                    self.log(f"Failed to split {file_path}: {e}")
                self.progress.setValue(idx)

        self.log("Processing complete.")
        QMessageBox.information(self, "Done", "Processing finished.")

    def merge_files(self, files, merged_path):
        """
        Merge multiple files into one workbook.
        """
        dfs = []
        colsets = []
        sheetmaps = {}
        for f in files:
            try:
                sheets = read_all_sheets(f)
                if not sheets: continue
                # prefer first sheet for column comparison
                first_sheet_name = list(sheets.keys())[0]
                df = sheets[first_sheet_name]
                dfs.append((f, df))
                colsets.append(tuple(df.columns))
                sheetmaps[f] = sheets
            except Exception as e:
                self.log(f"Skipping {f} during merge: {e}")

        # if all column sets identical, concat
        if len(colsets) >= 1 and all(cs == colsets[0] for cs in colsets):
            merged_df = pd.concat([df for _, df in dfs], ignore_index=True)
            save_df_to_excel(merged_df, merged_path, sheet_name='Merged')
        else:
            # create workbook with each file as sheet
            wb = Workbook()
            wb.remove(wb.active)
            for f, sheets in sheetmaps.items():
                short = os.path.splitext(os.path.basename(f))[0][:25]
                for sheetname, df in sheets.items():
                    title = f"{short}__{sheetname}"[:31]
                    ws = wb.create_sheet(title=title)
                    ws.append(list(df.columns))
                    for _, r in df.iterrows():
                        ws.append([r.get(c) for c in df.columns])
            wb.save(merged_path)

        apply_openpyxl_autofit_and_theme(merged_path)

    def log(self, text):
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        self.log_text.append(f"[{now}] {text}")
        cursor = self.log_text.textCursor()
        self.log_text.moveCursor(cursor.End)

def main():
    app = QApplication(sys.argv)
    w = ExcelMergeSplitWindow()
    w.show()
    sys.exit(app.exec_())

if __name__ == "__main__":
    main()

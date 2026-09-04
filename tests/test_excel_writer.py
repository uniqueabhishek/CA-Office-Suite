"""
Tests for core.excel_writer.

The column-width assertions are regression tests: the width clamp previously
read min(max(50, max_len + 2), 100), which forced every column in every
generated workbook to at least 50 characters wide.
"""

import os

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from config.constants import MAX_COLUMN_WIDTH, MIN_COLUMN_WIDTH
from core.excel_writer import _fitted_width, apply_formatting_to_workbook, save_df_to_excel


class TestFittedWidth:
    """Column width clamping."""

    def test_short_content_gets_minimum_width(self):
        assert _fitted_width(1) == MIN_COLUMN_WIDTH

    def test_long_content_is_capped(self):
        assert _fitted_width(500) == MAX_COLUMN_WIDTH

    def test_mid_range_content_tracks_length(self):
        assert _fitted_width(20) == 22

    def test_width_never_leaves_the_clamp_range(self):
        for length in range(0, 200):
            assert MIN_COLUMN_WIDTH <= _fitted_width(length) <= MAX_COLUMN_WIDTH


class TestSaveDfToExcel:
    """Writing a DataFrame to a new workbook."""

    def test_short_columns_are_not_forced_to_fifty(self, simple_df, tmp_path):
        out = os.path.join(tmp_path, "out.xlsx")
        save_df_to_excel(simple_df, out)

        ws = load_workbook(out).active
        widths = [d.width for d in ws.column_dimensions.values()]
        assert widths, "expected column widths to be set"
        assert all(w == MIN_COLUMN_WIDTH for w in widths), widths

    def test_wide_column_is_capped(self, tmp_path):
        df = pd.DataFrame({"Long": ["x" * 300]})
        out = os.path.join(tmp_path, "wide.xlsx")
        save_df_to_excel(df, out)

        ws = load_workbook(out).active
        assert ws.column_dimensions["A"].width == MAX_COLUMN_WIDTH

    def test_writes_header_and_rows(self, simple_df, tmp_path):
        out = os.path.join(tmp_path, "rows.xlsx")
        save_df_to_excel(simple_df, out, sheet_name="Data")

        ws = load_workbook(out).active
        assert ws.title == "Data"
        assert [c.value for c in ws[1]] == ["A", "B"]
        assert ws.max_row == 3  # header + 2 rows

    def test_header_row_is_bold(self, simple_df, tmp_path):
        out = os.path.join(tmp_path, "bold.xlsx")
        save_df_to_excel(simple_df, out)

        ws = load_workbook(out).active
        assert all(c.font.bold for c in ws[1])

    def test_number_format_map_is_applied(self, tmp_path):
        df = pd.DataFrame({"Amount": [1.5, 2.25]})
        out = os.path.join(tmp_path, "fmt.xlsx")
        save_df_to_excel(df, out, number_format_map={"Amount": "#,##0.00"})

        ws = load_workbook(out).active
        assert ws["A2"].number_format == "#,##0.00"

    def test_creates_missing_output_directory(self, simple_df, tmp_path):
        out = os.path.join(tmp_path, "nested", "deeper", "out.xlsx")
        save_df_to_excel(simple_df, out)
        assert os.path.exists(out)


class TestApplyFormattingToWorkbook:
    """In-memory formatting of an existing workbook."""

    @pytest.fixture
    def workbook(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Name", "Amount"])
        ws.append(["Alice", 100])
        return wb

    def test_short_columns_are_not_forced_to_fifty(self, workbook):
        apply_formatting_to_workbook(workbook, apply_autofit=True)

        ws = workbook.active
        widths = [d.width for d in ws.column_dimensions.values()]
        assert widths
        assert all(w == MIN_COLUMN_WIDTH for w in widths), widths

    def test_autofit_off_leaves_widths_untouched(self, workbook):
        apply_formatting_to_workbook(workbook, apply_autofit=False)
        assert not workbook.active.column_dimensions

    def test_number_format_matched_by_header_name(self, workbook):
        apply_formatting_to_workbook(workbook, number_format_map={"Amount": "#,##0"})
        assert workbook.active["B2"].number_format == "#,##0"

    def test_returns_the_same_workbook(self, workbook):
        assert apply_formatting_to_workbook(workbook) is workbook

    def test_theme_freezes_the_header_row(self, workbook):
        apply_formatting_to_workbook(workbook, apply_theme=True)
        assert workbook.active.freeze_panes == "A2"

    def test_theme_adds_a_filter_over_the_used_range(self, workbook):
        apply_formatting_to_workbook(workbook, apply_theme=True)
        assert workbook.active.auto_filter.ref == "A1:B2"

    def test_theme_off_leaves_the_sheet_alone(self, workbook):
        """apply_theme used to be accepted and then ignored entirely."""
        apply_formatting_to_workbook(workbook, apply_theme=False)
        assert workbook.active.freeze_panes is None
        assert workbook.active.auto_filter.ref is None

    def test_empty_sheet_does_not_raise(self):
        wb = Workbook()
        apply_formatting_to_workbook(wb, apply_autofit=True)

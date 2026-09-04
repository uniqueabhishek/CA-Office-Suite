"""
Tests for core.merge_logic - the merge behaviour shared by both front-ends.
"""

import os

import pandas as pd
import pytest
from openpyxl import load_workbook

from config.constants import MAX_SHEET_NAME_LENGTH
from core.merge_logic import merge_files_logic


@pytest.fixture
def same_shape_files(tmp_path):
    """Two workbooks with identical columns."""
    paths = []
    for i, rows in enumerate([[('a', 1), ('b', 2)], [('c', 3)]]):
        p = os.path.join(tmp_path, f'same_{i}.xlsx')
        pd.DataFrame(rows, columns=['Name', 'Amount']).to_excel(p, index=False)
        paths.append(p)
    return paths


@pytest.fixture
def different_shape_files(tmp_path):
    """Two workbooks with different columns."""
    p1 = os.path.join(tmp_path, 'first.xlsx')
    p2 = os.path.join(tmp_path, 'second.xlsx')
    pd.DataFrame({'Name': ['a']}).to_excel(p1, index=False)
    pd.DataFrame({'Totally': ['x'], 'Different': ['y']}).to_excel(p2, index=False)
    return [p1, p2]


class TestMergeFilesLogic:

    def test_matching_columns_concatenate_into_one_sheet(self, same_shape_files, tmp_path):
        out = os.path.join(tmp_path, 'merged.xlsx')
        merge_files_logic(same_shape_files, out)

        wb = load_workbook(out)
        assert wb.sheetnames == ['Merged']
        assert wb['Merged'].max_row == 4  # header + 3 data rows

    def test_differing_columns_get_one_sheet_each(self, different_shape_files, tmp_path):
        out = os.path.join(tmp_path, 'merged.xlsx')
        merge_files_logic(different_shape_files, out)

        wb = load_workbook(out)
        assert len(wb.sheetnames) == 2
        assert 'Merged' not in wb.sheetnames

    def test_sheet_names_respect_the_excel_length_limit(self, tmp_path):
        long_name = 'a_very_long_source_file_name_indeed'
        p1 = os.path.join(tmp_path, f'{long_name}.xlsx')
        p2 = os.path.join(tmp_path, 'other.xlsx')
        pd.DataFrame({'A': [1]}).to_excel(p1, index=False)
        pd.DataFrame({'B': [1]}).to_excel(p2, index=False)

        out = os.path.join(tmp_path, 'merged.xlsx')
        merge_files_logic([p1, p2], out)

        wb = load_workbook(out)
        assert all(len(name) <= MAX_SHEET_NAME_LENGTH for name in wb.sheetnames)

    def test_unreadable_file_is_skipped(self, same_shape_files, tmp_path):
        broken = os.path.join(tmp_path, 'broken.xlsx')
        with open(broken, 'w', encoding='utf-8') as fh:
            fh.write('this is not a workbook')

        out = os.path.join(tmp_path, 'merged.xlsx')
        merge_files_logic(same_shape_files + [broken], out)

        assert os.path.exists(out)
        assert load_workbook(out)['Merged'].max_row == 4

    def test_output_is_formatted(self, same_shape_files, tmp_path):
        out = os.path.join(tmp_path, 'merged.xlsx')
        merge_files_logic(same_shape_files, out)

        ws = load_workbook(out)['Merged']
        assert all(c.font.bold for c in ws[1])
        assert ws.column_dimensions['A'].width

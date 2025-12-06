import pandas as pd
import os
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from core.excel_utils import read_all_sheets
from core.excel_writer import save_df_to_excel

def merge_files_logic(files, merged_path):
    """
    Merge multiple files into one workbook.
    Identical behavior to Desktop App logic including column-comparison for concat.
    """
    dfs = []
    colsets = []
    sheetmaps = {}

    # Reading Phase
    for f in files:
        try:
            sheets = read_all_sheets(f)
            if not sheets:
                continue
            # prefer first sheet for column comparison
            first_sheet_name = list(sheets.keys())[0]
            df = sheets[first_sheet_name]
            dfs.append((f, df))
            colsets.append(tuple(df.columns))
            sheetmaps[f] = sheets
        except Exception as e:
            # We catch exception but in web context we might want to log it
            print(f"Skipping {f} during merge: {e}")

    # Decision Phase: Concatenate or Separate Sheets?
    # if all column sets identical, concat
    if len(colsets) >= 1 and all(cs == colsets[0] for cs in colsets):
        merged_df = pd.concat([df for _, df in dfs], ignore_index=True)
        # Using save_df_to_excel from core, which handles writer logic
        save_df_to_excel(merged_df, merged_path, sheet_name='Merged')
    else:
        # create workbook with each file as sheet
        wb = Workbook()
        if wb.active is not None:
            wb.remove(wb.active)
        for f, sheets in sheetmaps.items():
            short = os.path.splitext(os.path.basename(f))[0][:25]
            for sheetname, df in sheets.items():
                title = f"{short}__{sheetname}"[:31]
                ws = wb.create_sheet(title=title)
                # Use fast dataframe_to_rows instead of slow iterrows
                for row in dataframe_to_rows(df, index=False, header=True):
                    ws.append(row)
        wb.save(merged_path)

    # Note: Desktop app calls apply_openpyxl_autofit_and_theme here.
    # We should likely include that if we want true parity.
    from core.excel_writer import apply_formatting_to_workbook

    # Re-open safely to apply styles
    from openpyxl import load_workbook
    wb = load_workbook(merged_path)
    apply_formatting_to_workbook(wb, apply_autofit=True) # Defaults from desktop
    wb.save(merged_path)

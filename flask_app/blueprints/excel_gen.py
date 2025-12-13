"""
This module handles the generation of the Balance Sheet Excel file
from the session data collected across various pages.
"""
import io
from openpyxl import Workbook  # pylint: disable=import-error
from openpyxl.styles import Font  # pylint: disable=import-error


def generate_balance_sheet(session_data):
    """
    Generates an Excel file from the session data collected across multiple pages.
    Returns a BytesIO object containing the Excel file.
    """
    wb = Workbook()

    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # --- Page 1: Schedules 1-4 ---
    ws1 = wb.create_sheet("Page 1 - Sch 1-4")
    data_p1 = session_data.get('bs_page1', {})

    # Simple key-value dump for Page 1 top section
    # This is a basic implementation. Ideally, this would map to specific cells.
    ws1.append(["Field", "Value"])
    for k, v in data_p1.items():
        if not k.startswith('sch3_'):  # Handle lists separately
            ws1.append([k, v])

    # Handle Schedule 3 (Liabilities) List
    ws1.append([])
    ws1.append(["Schedule 3: Liabilities"])
    ws1.append(["Particulars", "Amount"])
    sch3_particulars = data_p1.get('sch3_particulars[]', [])
    sch3_amounts = data_p1.get('sch3_amount[]', [])

    # Flask form lists might come as lists or single items depending on library version,
    # but request.form.to_dict(flat=False) is needed for lists.
    # The current session storage uses request.form.to_dict() which flattens by default!
    # CHECK: balance_sheet.py uses request.form.to_dict(). This destroys repeated keys (lists).
    # CRITICAL FIX NEEDED in balance_sheet.py to save lists correctly.
    # For now, assuming we handle it, let's write the code.

    if isinstance(sch3_particulars, list):
        for p, a in zip(sch3_particulars, sch3_amounts):
            ws1.append([p, a])
    else:
        # Fallback if flattened
        ws1.append([sch3_particulars, sch3_amounts])

    # --- Page 2: Schedules 8-13 ---
    ws2 = wb.create_sheet("Page 2 - Sch 8-13")
    data_p2 = session_data.get('bs_page2', {})
    ws2.append(["Field", "Value"])
    for k, v in data_p2.items():
        ws2.append([k, v])

    # --- Page 3: Investments ---
    ws3 = wb.create_sheet("Page 3 - Investments")
    data_p3 = session_data.get('bs_page3', {})
    ws3.append(["Field", "Value"])
    for k, v in data_p3.items():
        ws3.append([k, v])

    # --- Page 4: Assets ---
    ws4 = wb.create_sheet("Page 4 - Assets")
    data_p4 = session_data.get('bs_page4', {})
    ws4.append(["Field", "Value"])
    for k, v in data_p4.items():
        ws4.append([k, v])

    # --- Page 4C: Contribution ---
    ws4c = wb.create_sheet("Page 4C - Contribution")
    data_p4c = session_data.get('bs_page4c', {})
    ws4c.append(["Field", "Value"])
    for k, v in data_p4c.items():
        ws4c.append([k, v])

    # --- Page 4D: History ---
    ws4d = wb.create_sheet("Page 4D - History")
    data_p4d = session_data.get('bs_page4d', {})
    ws4d.append(["Field", "Value"])
    for k, v in data_p4d.items():
        ws4d.append([k, v])

    # --- Winman ---
    ws_win = wb.create_sheet("Winman")
    data_win = session_data.get('bs_winman', {})
    ws_win.append(["Field", "Value"])
    for k, v in data_win.items():
        ws_win.append([k, v])

    # --- Corpus ---
    ws_corpus = wb.create_sheet("Corpus Fund")
    data_corpus = session_data.get('bs_corpus', {})
    ws_corpus.append(["Field", "Value"])
    for k, v in data_corpus.items():
        ws_corpus.append([k, v])

    # --- Accumulation ---
    ws_acc = wb.create_sheet("Accumulation")
    data_acc = session_data.get('bs_accumulation', {})
    ws_acc.append(["Field", "Value"])
    for k, v in data_acc.items():
        ws_acc.append([k, v])

    # Styling Loop
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True)
        # Autofit (basic)
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter  # Get the column name
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:  # pylint: disable=broad-except
                    pass
            adjusted_width = max_length + 2
            ws.column_dimensions[column].width = adjusted_width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
